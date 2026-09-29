# -*- coding: utf-8 -*-
"""S7 导出：CSV / Markdown / XLSX，以及飞书多维表格的推送载荷。"""
from __future__ import annotations

import csv
import json
from pathlib import Path

from .s5_vision import ANALYSIS_FIELDS
from .s6_assemble import COLUMNS
from .utils import ensure_dir, log, run_cmd, save_json, sec_to_short, Timer


# ---------------------------------------------------------------- CSV

def to_csv(data: dict, out_path: Path) -> Path:
    """UTF-8 BOM，Excel / WPS / 飞书导入都不乱码。"""
    ensure_dir(out_path.parent)
    cols = [c for c, _ in COLUMNS]
    with out_path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for r in data["rows"]:
            w.writerow([_cell(r.get(_key_for(c))) for c in cols])
    return out_path


def _key_for(col: str) -> str:
    for c, k in COLUMNS:
        if c == col:
            return k
    return col


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:.3f}".rstrip("0").rstrip(".")
    return str(v)


# ---------------------------------------------------------------- Markdown

def to_markdown(data: dict, out_path: Path, meta: dict) -> Path:
    st = data["stats"]
    lines = [
        f"# 拉片表 · {meta.get('title') or '未命名'}",
        "",
        f"- **来源**：{meta.get('source', '')}",
        f"- **时长**：{st['total_duration']:.1f}s（{sec_to_short(st['total_duration'])}）",
        f"- **镜头数**：{st['shot_count']} · 平均镜头时长 {st['avg_shot_duration']}s",
        f"- **已完成画面分析**：{st['analyzed']}/{st['shot_count']} 镜",
        f"- **含台词镜头**：{st['with_dialogue']} 镜",
        "",
        "> 「人工分析」列留空，供人工补充。",
        "",
        "---",
        "",
        "## 一、拉片总表",
        "",
    ]

    cols = [c for c, _ in COLUMNS if c != "关键帧"]
    lines.append("| " + " | ".join(cols) + " |")
    lines.append("|" + "|".join(["---"] * len(cols)) + "|")
    for r in data["rows"]:
        cells = []
        for c in cols:
            v = _cell(r.get(_key_for(c)))
            v = v.replace("|", "\\|").replace("\n", " ")
            if len(v) > 90:
                v = v[:88] + "…"
            cells.append(v or " ")
        lines.append("| " + " | ".join(cells) + " |")

    lines += ["", "---", "", "## 二、逐镜详表（含关键帧）", ""]
    for r in data["rows"]:
        lines += [
            f"### 镜 {r['shot_id']} · {r['start_short']}–{r['end_short']}（{r['duration']}s）",
            "",
            f"![镜{r['shot_id']}]({_first_frame(r)})",
            "",
        ]
        for key, cn in ANALYSIS_FIELDS:
            if key == "visual_content":
                continue
            val = _cell(r.get(key))
            if val:
                lines.append(f"- **{cn}**：{val}")
        if r.get("dialogue"):
            lines.append(f"- **台词**：{r['dialogue']}")
        lines += [
            f"- **画面内容**：{_cell(r.get('visual_content')) or '（待分析）'}",
            "- **人工分析**：",
            "",
            "---",
            "",
        ]

    out_path.write_text("\n".join(lines), encoding="utf-8")
    return out_path


def _first_frame(r: dict) -> str:
    kf = r.get("keyframes") or ""
    return kf.split(" | ")[0] if kf else ""


def to_markdown_lite(data: dict, out_path: Path, meta: dict) -> Path:
    """不带图片的轻量版，适合直接粘进飞书文档 / 知识库。"""
    st = data["stats"]
    lines = [
        f"# 拉片表 · {meta.get('title') or '未命名'}",
        "",
        f"来源：{meta.get('source', '')}｜时长 {st['total_duration']:.1f}s｜"
        f"{st['shot_count']} 镜｜平均 {st['avg_shot_duration']}s／镜",
        "",
    ]
    for r in data["rows"]:
        lines += [
            f"## 镜{r['shot_id']}｜{r['start_short']}–{r['end_short']}｜{r['duration']}s",
            "",
            f"- 景别：{_cell(r.get('shot_size')) or '—'}｜运镜：{_cell(r.get('camera_move')) or '—'}",
            f"- 画面：{_cell(r.get('visual_content')) or '（待分析）'}",
            f"- 台词：{_cell(r.get('dialogue')) or '—'}",
            f"- 情绪：{_cell(r.get('emotion')) or '—'}",
            f"- 表达核心：{_cell(r.get('core_message')) or '—'}",
            f"- 用户体验路径：{_cell(r.get('ux_path')) or '—'}",
            f"- 视听分工：{_cell(r.get('av_division')) or '—'}",
            "- 人工分析：",
            "",
        ]
    out_path.write_text("\n".join(lines), encoding="utf-8")
    return out_path


# ---------------------------------------------------------------- XLSX

def to_xlsx(data: dict, out_path: Path) -> Path | None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        log("未安装 openpyxl，跳过 xlsx 导出", "warn")
        return None

    ensure_dir(out_path.parent)
    wb = Workbook()
    ws = wb.active
    ws.title = "拉片表"

    cols = [c for c, _ in COLUMNS]
    widths = {
        "镜号": 8, "起始": 10, "结束": 10, "时长(s)": 9, "景别": 10, "运镜": 10,
        "画面内容": 46, "关键帧": 30, "台词": 40, "情绪": 14,
        "表达核心": 36, "用户体验路径": 36, "视听分工": 36, "人工分析": 30,
    }
    head_fill = PatternFill("solid", fgColor="1F4E79")
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 18)
    ws.freeze_panes = "A2"

    for ri, r in enumerate(data["rows"], start=2):
        for ci, c in enumerate(cols, start=1):
            v = r.get(_key_for(c))
            if c == "关键帧":
                v = _first_frame(r)
            cell = ws.cell(row=ri, column=ci, value=_cell(v) if not isinstance(v, (int, float)) else v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if ri % 2 == 0:
            for ci in range(1, len(cols) + 1):
                ws.cell(row=ri, column=ci).fill = PatternFill("solid", fgColor="F2F7FB")

    wb.save(out_path)
    return out_path


# ---------------------------------------------------------------- 飞书载荷

def to_feishu_payload(data: dict, meta: dict, out_path: Path) -> Path:
    """生成飞书多维表格批量写入载荷，供 lark-cli / skill 直接推送。"""
    records = []
    for r in data["rows"]:
        fields = {
            "镜号": r["shot_id"],
            "起始": r["start_tc"],
            "结束": r["end_tc"],
            "时长(s)": r["duration"],
            "景别": r.get("shot_size", ""),
            "运镜": r.get("camera_move", ""),
            "画面内容": r.get("visual_content", ""),
            "台词": r.get("dialogue", ""),
            "情绪": r.get("emotion", ""),
            "表达核心": r.get("core_message", ""),
            "用户体验路径": r.get("ux_path", ""),
            "视听分工": r.get("av_division", ""),
            "人工分析": "",
        }
        records.append({"fields": fields})

    payload = {
        "title": meta.get("title"),
        "source": meta.get("source"),
        "stats": data["stats"],
        "columns": [c for c, _ in COLUMNS],
        "records": records,
    }
    save_json(payload, out_path)
    return out_path


# ---------------------------------------------------------------- 入口

def run(work_dir: Path, cfg: dict, data: dict, meta: dict) -> dict:
    """执行 S7，返回产出文件路径表。"""
    exp = cfg.get("export", {})
    out = ensure_dir(work_dir / "export")
    title = meta.get("title") or "拉片表"
    safe = "".join(ch for ch in title if ch not in '/\\:*?"<>|')[:40] or "拉片表"

    made: dict[str, str] = {}
    with Timer("S7 导出拉片表"):
        if exp.get("csv", True):
            p = to_csv(data, out / f"{safe}_拉片表.csv")
            made["csv"] = str(p)
        if exp.get("markdown", True):
            made["markdown"] = str(to_markdown(data, out / f"{safe}_拉片表.md", meta))
            made["markdown_lite"] = str(
                to_markdown_lite(data, out / f"{safe}_拉片表_简洁版.md", meta)
            )
        if exp.get("xlsx", False):
            p = to_xlsx(data, out / f"{safe}_拉片表.xlsx")
            if p:
                made["xlsx"] = str(p)

        made["feishu_payload"] = str(
            to_feishu_payload(data, meta, work_dir / "feishu_payload.json")
        )

    for k, v in made.items():
        if k != "feishu_payload":
            log(f"已生成 {k}：{Path(v).name}", "ok")
    return made
